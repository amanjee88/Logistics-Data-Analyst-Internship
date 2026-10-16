"""Build the formatted Week 4 Word report from executed model outputs.
Run from repository root: python src/build_week4_report.py
"""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
OUT = REPORTS / "Week_4_Predictive_Modeling_and_Optimization.docx"
ANALYSIS = ROOT / "data" / "analysis"
FIG = REPORTS / "figures"
NAVY = "17365D"
BLUE = "2F5597"
TEAL = "168A8A"
PALE = "EAF1F8"
PALE_TEAL = "E6F4F3"
LIGHT = "F4F6F8"
MID = "5B6573"
DARK = "202A35"
WHITE = "FFFFFF"
BORDER = "D9E1F2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tcMar = tc_pr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tc_pr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['val', 'sz', 'space', 'color']:
                if key in kwargs[edge]:
                    element.set(qn('w:{}'.format(key)), str(kwargs[edge][key]))


def set_font(run, name='Arial', size=10.2, color=DARK, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = field
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = '1'
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.append(begin); run._r.append(instr); run._r.append(separate); run._r.append(text); run._r.append(end)
    run.font.name = 'Arial'; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(MID)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.67)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.29)
    sec.footer_distance = Inches(0.3)
    sec.different_first_page_header_footer = True

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'; normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    normal.font.size = Pt(10.2); normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in [('Title', 31, NAVY), ('Subtitle', 16, TEAL), ('Heading 1', 17, NAVY), ('Heading 2', 12.5, BLUE), ('Heading 3', 10.8, TEAL)]:
        st = doc.styles[style_name]
        st.font.name = 'Arial'; st._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        st.font.size = Pt(size); st.font.bold = style_name != 'Subtitle'
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(12 if style_name != 'Title' else 0)
        st.paragraph_format.space_after = Pt(5)
    # Running header for all pages except cover.
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(2)
    set_font(hp.add_run('LOGISTICS DATA ANALYST  /  WEEK 4'), size=8, color=NAVY, bold=True)
    # Light bottom border under header.
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '5'); bottom.set(qn('w:space'), '3'); bottom.set(qn('w:color'), BORDER)
    pBdr.append(bottom); pPr.append(pBdr)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(2)
    fp.paragraph_format.space_after = Pt(0)
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(fp.add_run('Synthetic teaching scenario  •  Not operational data'), size=7.5, color=MID, italic=True)
    set_font(fp.add_run(' ' * 4 + 'Page '), size=8, color=MID)
    add_field(fp, 'PAGE')


def add_para(doc, text='', bold_lead=None, italic=False, color=DARK, size=10.2, align=None, after=5.5):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead), size=size, color=color, bold=True)
        set_font(p.add_run(text[len(bold_lead):]), size=size, color=color, italic=italic)
    else:
        set_font(p.add_run(text), size=size, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.left_indent = Inches(0.22 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(text), size=9.7)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), size=9.6)
    return p


def add_table(doc, headers, rows, widths=None, font_size=8.3, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        for i, w in enumerate(widths):
            table.columns[i].width = Inches(w)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, val in enumerate(headers):
        cell = hdr.cells[i]
        if widths: cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, header_fill); set_cell_margins(cell)
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(str(val)), size=font_size, color=WHITE, bold=True)
        set_cell_border(cell, bottom={'val':'single','sz':'8','color':header_fill})
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, val in enumerate(row):
            cell = cells[i]
            if widths: cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=72, bottom=72, start=95, end=95)
            if ridx % 2 == 1: set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.02
            set_font(p.add_run(str(val)), size=font_size, color=DARK)
            set_cell_border(cell, bottom={'val':'single','sz':'3','color':BORDER})
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, title, body, fill=PALE_TEAL, title_color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill); set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    set_cell_border(cell, left={'val':'single','sz':'20','color':title_color})
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(title), size=10.3, color=title_color, bold=True)
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.08
    set_font(p2.add_run(body), size=9.2, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, filename, caption, width=6.2):
    path = FIG / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(7)
    set_font(cp.add_run(caption), size=8.2, color=MID, italic=True)


def add_side_by_side_figures(doc, left_file, left_caption, right_file, right_caption, width=3.05):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for i, (fn, cap) in enumerate([(left_file, left_caption), (right_file, right_caption)]):
        cell = table.cell(0, i); cell.width = Inches(width); set_cell_margins(cell, top=35, bottom=35, start=45, end=45)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(FIG / fn), width=Inches(width - .1))
        cp = cell.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(2)
        set_font(cp.add_run(cap), size=7.6, color=MID, italic=True)
        for edge in ['top','left','bottom','right','insideH','insideV']:
            set_cell_border(cell, **{edge:{'val':'nil'}})
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_num_cover(doc):
    pass


def main():
    REPORTS.mkdir(exist_ok=True)
    metrics = pd.read_csv(ANALYSIS / 'week4_model_metrics.csv')
    cv = pd.read_csv(ANALYSIS / 'week4_cv_metrics.csv')
    summary = json.loads((ANALYSIS / 'week4_modeling_summary.json').read_text(encoding='utf-8'))
    opt = json.loads((ANALYSIS / 'week4_optimization_summary.json').read_text(encoding='utf-8'))
    top = pd.read_csv(ANALYSIS / 'week4_feature_importance.csv').head(7)
    region_mae = pd.read_csv(ANALYSIS / 'week4_test_error_by_region.csv')
    quality = pd.read_csv(ANALYSIS / 'week4_data_quality_checks.csv')

    lr = metrics.loc[metrics.model == 'Linear regression'].iloc[0]
    rf = metrics.loc[metrics.model == 'Tuned random forest'].iloc[0]
    dummy = metrics.loc[metrics.model == 'Median baseline'].iloc[0]
    best_model = summary['selected_model']
    opt_cost_reduction = (opt['mean_estimated_cost_current_inr'] - opt['mean_estimated_cost_recommended_inr']) / opt['mean_estimated_cost_current_inr'] * 100

    doc = Document()
    style_document(doc)
    sec = doc.sections[0]
    # Cover page content.
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(9)
    set_font(p.add_run('WEEK 4  |  TECHNICAL PROJECT REPORT'), size=10, color=TEAL, bold=True)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run('Predictive Modeling\nand Optimization in\nLogistics Systems'), size=29, color=NAVY, bold=True)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(18)
    set_font(p.add_run('Delivery-time forecasting, model validation, and capacity-constrained what-if decisions'), size=15, color=TEAL)
    line = doc.add_paragraph(); line.paragraph_format.space_after = Pt(15)
    set_font(line.add_run('______________________________________________'), size=12, color=TEAL)
    add_para(doc, 'Prepared for: YuvaIntern — Logistics Data Analyst Intern', size=11.2, color=NAVY, after=7)
    add_para(doc, 'Prepared by: Aman Yadav', size=11.2, color=NAVY, after=7)
    add_para(doc, 'Project period: 16 October 2026', size=10.2, color=MID, after=18)
    add_callout(doc, 'DATA STATUS', 'This report uses 1,200 deterministically generated synthetic shipment records. The data, model metrics, estimated costs, and what-if recommendations are demonstration outputs only—not records or measured results from a real logistics operator.', fill=PALE)
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(18)
    set_font(p.add_run('A reproducible Python project with source code, model outputs, automated checks, figures, and a documented evaluation approach.'), size=9.4, color=MID, italic=True)

    doc.add_page_break()
    doc.add_heading('1. Executive Summary', level=1)
    add_para(doc, f'This project develops a reproducible regression workflow to estimate delivery time in hours from shipment and dispatch-context features. It creates a synthetic dataset of 1,200 shipments, reserves 20% (240 rows) as an untouched holdout set, and compares a median baseline, Linear Regression, and a hyperparameter-tuned Random Forest. Model choice is made from training-set five-fold cross-validation MAE, not from the holdout test results.')
    add_para(doc, f'Linear Regression was selected because it achieved the lowest five-fold cross-validation MAE among the predictive candidates ({cv.loc[cv.model == "Linear regression", "mean_cv_mae_hours"].iloc[0]:.3f} hours, standard deviation {cv.loc[cv.model == "Linear regression", "std_cv_mae_hours"].iloc[0]:.3f}). On the untouched test set, it achieved MAE {lr.MAE_hours:.3f} hours, RMSE {lr.RMSE_hours:.3f} hours, and R-squared {lr.R2:.3f}. These scores are consistent with a structured simulation in which the target is generated from known factors; they must not be presented as expected performance on real carrier data.')
    add_para(doc, f'A capacity-constrained, per-shipment what-if heuristic was also demonstrated on the 240 holdout shipments. It changed the suggested vehicle for {opt["shipments_with_vehicle_change"]} shipments ({opt["vehicle_change_share_pct"]:.1f}%). Under illustrative cost rates and a time value of ₹{opt["time_value_inr_per_hour"]:.0f} per predicted hour, mean estimated cost fell from ₹{opt["mean_estimated_cost_current_inr"]:,.2f} to ₹{opt["mean_estimated_cost_recommended_inr"]:,.2f} per shipment, while predicted time changed from {opt["mean_predicted_hours_current"]:.2f} to {opt["mean_predicted_hours_recommended"]:.2f} hours. This is a scenario exercise, not actual or validated business savings.')
    add_callout(doc, 'MAIN TAKEAWAY', 'A good analytical workflow separates training, model selection, holdout evaluation, and operational what-if analysis. Prediction quality does not itself prove that a dispatch policy will improve real-world performance.')

    doc.add_heading('2. Problem Definition and Objectives', level=1)
    add_para(doc, 'Business scenario: a fictional e-commerce logistics operator wants a dispatch-time estimate of delivery duration to help planners identify shipments that may exceed an illustrative 12-hour service threshold, compare candidate predictive models, and explore vehicle choices within payload limits.')
    add_para(doc, 'Target variable: delivery_time_hours, a positive numeric duration between dispatch and delivery. Predictors are fields assumed available at dispatch or estimated at that moment: distance, shipment payload, route stops, region, vehicle type, traffic/weather severity, warehouse load, priority, and day of week.')
    add_bullet(doc, 'Build a transparent benchmark and compare a linear model with a non-linear ensemble model.')
    add_bullet(doc, 'Evaluate with MAE, RMSE and R-squared, supported by five-fold cross-validation and an untouched test split.')
    add_bullet(doc, 'Use forecasts for exception triage and a capacity-constrained vehicle what-if, while explaining uncertainty and assumptions.')

    doc.add_page_break()
    doc.add_heading('3. Dataset, Features, and Assumptions', level=1)
    add_para(doc, 'The generator creates 1,200 rows with a fixed random seed (20261016). The dataset is reproducible, deliberately lightweight, and fully synthetic. No third-party raw shipment file is downloaded or treated as an empirical source for this task.')
    add_table(doc, ['Feature', 'Type', 'Use in model / definition'], [
        ['distance_km', 'Numeric', 'Estimated dispatch-to-destination distance; principal duration driver in the simulation.'],
        ['weight_kg; volume_m3', 'Numeric', 'Payload measurements used for transit modeling and vehicle-capacity checks.'],
        ['route_stops', 'Integer', 'Number of planned stops on the route.'],
        ['traffic_index', 'Ordinal 1–5', 'Synthetic traffic severity available at dispatch time.'],
        ['weather_severity', 'Ordinal 0–3', 'Synthetic weather disruption indicator.'],
        ['warehouse_load_pct', 'Numeric', 'Illustrative warehouse loading pressure (percent).'],
        ['vehicle_type', 'Category', 'Bike, Van, or Truck; used in time model and capacity what-if.'],
        ['origin_region', 'Category', 'North, South, East, West, or Central.'],
        ['priority; day_of_week', 'Category', 'Service priority and dispatch weekday.'],
        ['delivery_time_hours', 'Target', 'Simulated observed duration, in hours; excluded from predictors.'],
    ], widths=[1.45, 1.0, 4.3], font_size=8.1)
    add_para(doc, 'Simulated delivery time is generated using distance divided by a vehicle-specific assumed speed, adjusted by traffic and weather factors, plus stop, warehouse, regional, priority, and random-noise terms. This construction makes the test feasible and repeatable, but it also means model quality is partly a reflection of the chosen data-generating mechanism.', size=9.4, color=MID)
    doc.add_heading('Data checks and preparation', level=2)
    add_bullet(doc, '1,200 unique shipment IDs and positive target values are validated before modeling.')
    add_bullet(doc, 'There are no missing values in the generated modeling columns by design; preprocessing pipelines still use median imputation for numeric inputs and most-frequent imputation for categorical inputs as a defensive reusable pattern.')
    add_bullet(doc, 'Numeric predictors are scaled for Linear Regression; categorical predictors are one-hot encoded with unknown categories ignored. Random Forest uses imputation and one-hot encoding without numeric scaling.')
    add_bullet(doc, 'Estimated transport cost is not used as a model predictor. This avoids making the time forecast dependent on a synthetic cost formula and keeps the selected feature set focused on dispatch context.')

    doc.add_page_break()
    doc.add_heading('4. Modeling and Validation Framework', level=1)
    add_table(doc, ['Stage', 'Implementation', 'Reason'], [
        ['1. Split', '80% train (960 rows) / 20% test (240 rows); fixed seed.', 'Keeps an untouched holdout for final generalization check.'],
        ['2. Baseline', 'DummyRegressor predicting the training-target median.', 'Checks whether feature-based models improve on a naive estimate.'],
        ['3. Linear model', 'Pipeline with imputers, StandardScaler for numeric fields, one-hot encoding, and LinearRegression.', 'Interpretable benchmark; efficient for a compact dataset.'],
        ['4. Ensemble model', 'RandomForestRegressor; grid search over 120/220 trees, max depth None/18, leaf size 1/3; 3-fold grid CV.', 'Captures non-linearities and feature interactions.'],
        ['5. CV model choice', 'Five-fold shuffled KFold on training set; compare mean absolute error (MAE).', 'Select model based on training folds, not test scores.'],
        ['6. Holdout evaluation', 'MAE, RMSE and R-squared on 240 test rows.', 'Final model check on observations not used for fitting or model selection.'],
        ['7. What-if analysis', 'Use selected model for candidate vehicle predictions, subject to weight/volume capacities.', 'Demonstrates how estimates can inform decisions without claiming live optimization.'],
    ], widths=[1.2, 3.1, 2.45], font_size=7.8)
    add_para(doc, 'The five-fold comparison used shuffled folds on the training set (seed 20261016). The Random Forest grid search used its own three-fold cross-validation restricted to the same training split. The holdout set was not used to choose the final model. Since the exercise is simulated and cross-sectional, a future real project with time-dependent records should use a chronological or rolling-origin validation design.', size=9.2)
    doc.add_heading('Evaluation metrics', level=2)
    add_bullet(doc, 'MAE (hours): mean absolute difference between actual and predicted duration; directly interpretable in hours.')
    add_bullet(doc, 'RMSE (hours): square root of mean squared error; penalizes large misses more heavily than MAE.')
    add_bullet(doc, 'R-squared: relative fit against predicting the target mean; it can be negative when a model is worse than that baseline.')
    add_figure(doc, 'week4_model_comparison.png', 'Figure 1. Holdout MAE and RMSE for the median baseline, Linear Regression, and tuned Random Forest.', width=5.9)

    doc.add_page_break()
    doc.add_heading('5. Model Results and Selection', level=1)
    metric_rows = []
    for _, r in metrics.iterrows():
        metric_rows.append([r['model'], f'{r.MAE_hours:.3f}', f'{r.RMSE_hours:.3f}', f'{r.R2:.3f}'])
    add_table(doc, ['Model', 'MAE (h)', 'RMSE (h)', 'R²'], metric_rows, widths=[3.1, 1.05, 1.05, 1.0], font_size=8.5)
    add_para(doc, f'Five-fold training CV MAE was {cv.loc[cv.model == "Median baseline", "mean_cv_mae_hours"].iloc[0]:.3f} ± {cv.loc[cv.model == "Median baseline", "std_cv_mae_hours"].iloc[0]:.3f} hours for the median baseline, {cv.loc[cv.model == "Linear regression", "mean_cv_mae_hours"].iloc[0]:.3f} ± {cv.loc[cv.model == "Linear regression", "std_cv_mae_hours"].iloc[0]:.3f} hours for Linear Regression, and {cv.loc[cv.model == "Random forest (default; 150 trees)", "mean_cv_mae_hours"].iloc[0]:.3f} ± {cv.loc[cv.model == "Random forest (default; 150 trees)", "std_cv_mae_hours"].iloc[0]:.3f} hours for the default 150-tree Random Forest. The lower mean CV MAE supports selecting Linear Regression among the predictive candidates before viewing holdout metrics.')
    add_para(doc, f'The tuned Random Forest grid search selected {summary["best_random_forest_params"]["n_estimators"]} trees, max_depth={summary["best_random_forest_params"]["max_depth"]}, and min_samples_leaf={summary["best_random_forest_params"]["min_samples_leaf"]}; its best three-fold training-grid MAE was {summary["random_forest_best_grid_cv_mae_hours"]:.3f} hours. However, tuning does not guarantee that the more complex model will beat the linear model on held-out data. Here Linear Regression achieved the lower test MAE ({lr.MAE_hours:.3f} vs {rf.MAE_hours:.3f} hours).')
    add_callout(doc, 'INTERPRETATION CAUTION', f'The selected Linear Regression achieved R² = {lr.R2:.3f} on the synthetic holdout. This strong score is expected to some degree because the synthetic target was constructed from known features and simple relationships. It is not evidence that a similar score can be reached using real transport records.')
    doc.add_heading('Cross-validation summary', level=2)
    cv_rows = [[r['model'], str(int(r.cv_folds)), f'{r.mean_cv_mae_hours:.3f}', f'{r.std_cv_mae_hours:.3f}'] for _, r in cv.iterrows()]
    add_table(doc, ['Model', 'Folds', 'Mean CV MAE (h)', 'Std. dev. (h)'], cv_rows, widths=[3.0, .75, 1.35, 1.1], font_size=8.1)

    doc.add_page_break()
    doc.add_heading('6. Prediction Diagnostics', level=1)
    add_para(doc, 'Diagnostic plots are used together: the actual-versus-predicted plot checks whether estimates generally track observed values; the residual distribution checks the direction and spread of misses. Neither plot can establish operational reliability beyond this synthetic sample.')
    add_side_by_side_figures(doc, 'week4_actual_vs_predicted.png', 'Figure 2. Holdout actual vs predicted duration; dashed line is perfect agreement.', 'week4_residual_distribution.png', 'Figure 3. Holdout residuals (actual minus predicted).', width=3.05)
    doc.add_heading('Regional error check', level=2)
    rows = [[r.origin_region, int(r.n_test_shipments), f'{r.mae_hours:.3f}'] for _, r in region_mae.iterrows()]
    add_table(doc, ['Origin region', 'Holdout shipments', 'MAE (hours)'], rows, widths=[2.1, 1.65, 1.65], font_size=8.4)
    add_figure(doc, 'week4_test_error_by_region.png', 'Figure 4. Holdout MAE by origin region. Segment counts are modest; differences are descriptive only.', width=5.9)
    add_para(doc, 'The highest segment MAE in this particular test sample is for North (about {:.3f} hours) while Central is lowest (about {:.3f} hours). This is a prompt for additional validation, not a ranking of real regional operations.'.format(float(region_mae.loc[region_mae.origin_region == 'North', 'mae_hours'].iloc[0]), float(region_mae.loc[region_mae.origin_region == 'Central', 'mae_hours'].iloc[0])), size=9.2)

    heading7 = doc.add_heading('7. Feature Diagnostics and SLA Triage', level=1)
    heading7.paragraph_format.page_break_before = True
    top_rows = []
    for _, r in top.head(6).iterrows():
        nm = r.feature.replace('numeric__', '').replace('categorical__', '')
        top_rows.append([nm, f'{r.importance:.4f}', f'{r.importance*100:.1f}%'])
    add_table(doc, ['Random Forest feature', 'Importance', 'Share'], top_rows, widths=[3.2, 1.2, 1.0], font_size=8.2)
    add_figure(doc, 'week4_feature_importance.png', 'Figure 5. Top Random Forest impurity-based feature importances; they are not causal effects.', width=5.9)
    add_para(doc, 'Distance dominates this Random Forest importance diagnostic because the data-generation mechanism calculates duration largely from distance and assumed vehicle speed. Weather, weight, warehouse load, stops, and traffic appear much smaller in this design. Feature importance depends on the synthetic formula, correlations between inputs, and the model; it must not be interpreted as a causal ranking.')
    add_heading_para = doc.add_heading('Illustrative service-exception triage', level=2)
    add_para(doc, 'The analysis flags a shipment for review when its predicted delivery time exceeds 12 hours. This threshold is an example operating rule, not a learned probability and not a validated service-level agreement. A production workflow should use destination-specific promises, operating calendars, and a calibrated estimate of uncertainty where possible.')

    doc.add_page_break()
    doc.add_heading('8. Capacity-Constrained Vehicle What-if', level=1)
    add_para(doc, 'To demonstrate a decision layer, each holdout shipment was scored against vehicle candidates that pass hypothetical weight and volume limits. For each candidate, the model predicts delivery hours and an illustrative cost formula estimates dispatch cost. The chosen candidate minimizes a simple composite objective: estimated cost in INR + ₹85 × predicted delivery hours.')
    add_table(doc, ['Vehicle', 'Weight limit', 'Volume limit', 'Illustrative cost formula'], [
        ['Bike', '25 kg', '0.12 m³', '₹75 fixed + ₹5.80/km + payload terms'],
        ['Van', '900 kg', '6.0 m³', '₹220 fixed + ₹15.80/km + payload terms'],
        ['Truck', '5,000 kg', '30 m³', '₹500 fixed + ₹23.00/km + payload terms'],
    ], widths=[.85, 1.0, 1.0, 3.4], font_size=7.8)
    add_callout(doc, 'OBJECTIVE FUNCTION', 'Score = estimated vehicle cost + (time value × predicted delivery time). Time value is set to ₹85 per hour for this example. Capacity thresholds, fixed charges, per-kilometre rates, and time value are assumptions to demonstrate the workflow; they are not sourced tariffs.', fill=PALE)
    opt_rows = [
        ['Holdout shipments evaluated', str(opt['n_shipments'])],
        ['Suggested vehicle changes', f'{opt["shipments_with_vehicle_change"]} ({opt["vehicle_change_share_pct"]:.1f}%)'],
        ['Mean predicted time: current → recommended', f'{opt["mean_predicted_hours_current"]:.2f} h → {opt["mean_predicted_hours_recommended"]:.2f} h'],
        ['Mean estimated cost: current → recommended', f'₹{opt["mean_estimated_cost_current_inr"]:,.2f} → ₹{opt["mean_estimated_cost_recommended_inr"]:,.2f}'],
        ['Illustrative mean cost change', f'{opt_cost_reduction:.1f}% lower under assumed rates'],
        ['Mean composite objective improvement', f'{opt["mean_objective_score_improvement_pct"]:.1f}%'],
    ]
    add_table(doc, ['Scenario measure', 'Holdout what-if result'], opt_rows, widths=[2.9, 3.25], font_size=8.2)
    add_side_by_side_figures(doc, 'week4_optimization_time_comparison.png', 'Figure 6. Mean predicted time under current vs recommended vehicle choices.', 'week4_optimization_cost_comparison.png', 'Figure 7. Mean estimated cost under current vs recommended choices.', width=3.05)
    add_para(doc, 'In this synthetic test subset, cost decreases under the assumed rate card but mean predicted time increases slightly (around 0.052 hours). The composite objective improves because it values lower estimated cost more heavily under the chosen ₹85/hour trade-off. Decision-makers should review the cost-time trade-off, verify service constraints, and test different time values before considering any policy change.', size=9.2)

    doc.add_page_break()
    doc.add_heading('9. Recommendations for Logistics Operations', level=1)
    recommendations = [
        ('Use forecasts for exception triage, not automatic promises.', 'Prioritize shipments predicted above a validated service threshold for dispatcher review. Display estimates with confidence or error ranges once calibrated on real data.'),
        ('Add a service-aware objective.', 'Set business-approved weights for transportation cost, late-delivery penalties, overtime, emissions, and service promises. Test alternative time values rather than assuming ₹85/hour is universal.'),
        ('Respect vehicle capacity and route feasibility.', 'Before a recommendation is actionable, check payload, volume, vehicle availability, driver hours, depot allocation, customer time windows, road restrictions, and multi-stop routing.'),
        ('Move from independent assignments to fleet-level optimization.', 'The current heuristic selects a vehicle separately for each shipment; it does not schedule an entire fleet. Google OR-Tools can model vehicle routing with capacities and time windows [4], after valid distance/time matrices and constraints are available.'),
        ('Establish a real-data pilot and monitoring loop.', 'Use time-stamped dispatch snapshots, actual arrival events, promised times, route and traffic data. Validate chronologically, monitor MAE/RMSE and SLA exception rates by region/vehicle, and retrain only with documented evidence.'),
    ]
    for title, body in recommendations:
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.1
        set_font(p.add_run(title + ' '), size=9.7, color=NAVY, bold=True)
        set_font(p.add_run(body), size=9.5, color=DARK)
    doc.add_heading('Implementation roadmap', level=2)
    add_table(doc, ['Stage', 'Action', 'Gate to proceed'], [
        ['1. Define', 'Confirm SLA definition, decision users, feature availability time, and cost/service objective.', 'Business owner approves definitions and constraints.'],
        ['2. Data readiness', 'Collect appropriate dispatch snapshots, actual delivery events, costs, and route context; audit leakage and missingness.', 'Data quality and temporal alignment documented.'],
        ['3. Validate', 'Compare baseline and candidate models using chronological holdout and route/region diagnostics.', 'Errors meet an agreed business tolerance.'],
        ['4. Shadow pilot', 'Generate recommendations without changing dispatch; compare with actual outcomes and planner judgement.', 'Safety, feasibility and service impact reviewed.'],
        ['5. Controlled pilot', 'Test a limited deployment with human review and rollback criteria.', 'Measured operational benefit demonstrated before scale-up.'],
    ], widths=[1.0, 3.0, 2.3], font_size=7.7)

    doc.add_heading('10. Limitations and Responsible Use', level=1)
    for text in [
        'All 1,200 shipment records, weather/traffic indicators, delivery times, vehicle assignments, cost rates, and scenario outputs are synthetic. There are no real customer or courier records.',
        'The target is generated from a known formula, so high holdout performance is a demonstration of recovering simulated relationships; it is not independent evidence of real-world utility.',
        'Random train/test splitting is used for this randomly generated, cross-sectional teaching exercise. Real dispatch data usually require chronological or rolling validation to control time leakage and drift.',
        'The vehicle what-if uses simplified capacity limits, a fixed rate card, and an independent per-shipment score. It is not a route sequence optimizer, does not model shared fleet resources, and has no measured savings guarantee.',
        'A regression estimate above 12 hours is a rule-based exception flag, not a calibrated delay probability. It should not be presented as a service guarantee.'
    ]: add_bullet(doc, text)

    doc.add_page_break()
    doc.add_heading('11. Reproducibility and Code Structure', level=1)
    add_para(doc, 'Run commands from the repository root in a Python 3.10+ environment. The dependency list is in requirements.txt. The script uses a fixed random seed for repeatability; package versions can still affect small numerical differences, so preserve the environment for exact reproduction.')
    code = '''pip install -r requirements.txt\npython src/generate_week4_hypothetical_data.py\npython src/train_evaluate_optimize.py\npython tests/test_week4_modeling.py\npython src/build_week4_report.py'''
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.right_indent = Inches(0.15); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(7)
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F3F5F7'); pPr.append(shd)
    for i, line in enumerate(code.splitlines()):
        if i: p.add_run('\n')
        set_font(p.add_run(line), name='Consolas', size=8.8, color=NAVY)
    doc.add_heading('Core model-selection example', level=2)
    snippet = '''cv = KFold(n_splits=5, shuffle=True, random_state=20261016)\nmae = -cross_val_score(\n    pipeline, X_train, y_train, cv=cv,\n    scoring="neg_mean_absolute_error"\n)\n# Select using training CV MAE; assess once on held-out test data.\nmodel.fit(X_train, y_train)\npred = model.predict(X_test)'''
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.right_indent = Inches(0.15); p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F3F5F7'); pPr.append(shd)
    for i, line in enumerate(snippet.splitlines()):
        if i: p.add_run('\n')
        set_font(p.add_run(line), name='Consolas', size=8.4, color=NAVY)
    add_para(doc, 'Repository files include the synthetic data generator, model-training/evaluation/optimization script, model metrics, CV results, hyperparameter search output, feature diagnostics, holdout predictions, what-if vehicle recommendations, figure files, and automated integrity tests.', size=9.2)

    doc.add_heading('References', level=1)
    references = [
        '[1] Scikit-learn developers. Model evaluation: quantifying the quality of predictions. https://scikit-learn.org/stable/modules/model_evaluation.html',
        '[2] Scikit-learn developers. LinearRegression API documentation. https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.LinearRegression.html',
        '[3] Scikit-learn developers. RandomForestRegressor API documentation. https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.RandomForestRegressor.html',
        '[4] Google for Developers. Vehicle Routing Problem — OR-Tools. https://developers.google.com/optimization/routing/vrp',
    ]
    for ref in references:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.22); p.paragraph_format.first_line_indent = Inches(-0.22); p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run(ref), size=8.8, color=DARK)
    add_callout(doc, 'FINAL NOTE', 'This is a training and portfolio project. Validate data definitions, model reliability, operational feasibility, and business benefit with real, properly governed logistics data before any production use.', fill=PALE)

    doc.core_properties.title = 'Week 4 - Predictive Modeling and Optimization in Logistics Systems'
    doc.core_properties.subject = 'Synthetic logistics regression modeling and what-if vehicle assignment'
    doc.core_properties.author = 'Aman Yadav'
    doc.core_properties.keywords = 'logistics, Python, regression, Random Forest, model evaluation, optimization'
    doc.save(OUT)
    print(f'Wrote {OUT}')


if __name__ == '__main__':
    main()
